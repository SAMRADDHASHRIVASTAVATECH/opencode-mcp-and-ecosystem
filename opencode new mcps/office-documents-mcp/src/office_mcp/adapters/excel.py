"""Excel (XLSX) adapter over openpyxl + pandas analysis."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference, ScatterChart
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.defined_name import DefinedName

from office_mcp.config import SETTINGS
from office_mcp.errors import NotFoundError, ValidationError

CHARTS = {
    "bar": BarChart,
    "line": LineChart,
    "pie": PieChart,
    "scatter": ScatterChart,
}


def create_workbook(path: Path, sheets: list[str] | None = None, title: str | None = None) -> None:
    wb = Workbook()
    if title:
        wb.properties.title = title
    default = wb.active
    names = sheets or ["Sheet1"]
    default.title = names[0][:31]
    for name in names[1:]:
        wb.create_sheet(name[:31])
    wb.save(str(path))


def load(path: Path, data_only: bool = False, read_only: bool = False):
    return load_workbook(str(path), data_only=data_only, read_only=read_only)


def inspect(path: Path) -> dict[str, Any]:
    wb = load(path, read_only=True)
    try:
        sheets = []
        for name in wb.sheetnames:
            ws = wb[name]
            dims = ws.calculate_dimension() if hasattr(ws, "calculate_dimension") else None
            sheets.append(
                {
                    "name": name,
                    "dimensions": dims,
                    "max_row": ws.max_row,
                    "max_column": ws.max_column,
                }
            )
        defined = []
        # read_only workbooks may still expose defined_names
        try:
            for dn in wb.defined_names.values():
                defined.append({"name": dn.name, "attr_text": dn.attr_text})
        except Exception:
            pass
        return {
            "sheets": sheets,
            "sheetnames": list(wb.sheetnames),
            "defined_names": defined,
        }
    finally:
        wb.close()


def _sheet(wb, name: str | None):
    if name is None:
        return wb.active
    if name not in wb.sheetnames:
        raise NotFoundError(f"Sheet not found: {name}", {"sheets": wb.sheetnames})
    return wb[name]


def read_range(
    path: Path,
    sheet: str | None = None,
    range_addr: str | None = None,
    include_formulas: bool = True,
    data_only: bool = False,
    max_cells: int | None = None,
) -> dict[str, Any]:
    cap = max_cells or SETTINGS.max_cells
    wb = load(path, data_only=data_only)
    try:
        ws = _sheet(wb, sheet)
        if range_addr:
            min_col, min_row, max_col, max_row = range_boundaries(range_addr)
        else:
            min_row, min_col = 1, 1
            max_row, max_col = ws.max_row or 1, ws.max_column or 1
        n = (max_row - min_row + 1) * (max_col - min_col + 1)
        if n > cap:
            raise ValidationError(
                f"Range has {n} cells, exceeds cap {cap}",
                {"range": range_addr, "cells": n},
            )
        values = []
        formulas = []
        for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            values.append([_cell_value(c) for c in row])
            if include_formulas and not data_only:
                formulas.append([c.value if isinstance(c.value, str) and str(c.value).startswith("=") else None for c in row])
        return {
            "sheet": ws.title,
            "range": range_addr or ws.dimensions,
            "values": values,
            "formulas": formulas if include_formulas and not data_only else None,
            "rows": len(values),
            "cols": len(values[0]) if values else 0,
        }
    finally:
        wb.close()


def _cell_value(cell: Any) -> Any:
    v = cell.value
    if v is None:
        return None
    if hasattr(v, "isoformat"):
        try:
            return v.isoformat()
        except Exception:
            return str(v)
    if isinstance(v, (int, float, str, bool)):
        return v
    return str(v)


def write_cells(
    path: Path,
    sheet: str | None,
    cells: list[dict[str, Any]] | None = None,
    start: str | None = None,
    values: list[list[Any]] | None = None,
) -> dict[str, Any]:
    wb = load(path)
    ws = _sheet(wb, sheet)
    written = 0
    if cells:
        for item in cells:
            addr = item.get("cell") or item.get("address")
            if not addr:
                raise ValidationError("Each cell item needs 'cell'")
            ws[addr] = item.get("value")
            written += 1
    if values is not None:
        origin = start or "A1"
        col_letter, row = coordinate_from_string(origin)
        col = column_index_from_string(col_letter)
        for r_i, row_vals in enumerate(values):
            for c_i, val in enumerate(row_vals):
                ws.cell(row=row + r_i, column=col + c_i, value=val)
                written += 1
    wb.save(str(path))
    wb.close()
    return {"written": written, "sheet": ws.title, "path": str(path)}


def manage_sheets(path: Path, action: str, name: str | None = None, new_name: str | None = None, index: int | None = None) -> dict[str, Any]:
    wb = load(path)
    action = action.lower()
    if action == "list":
        names = list(wb.sheetnames)
        wb.close()
        return {"sheets": names}
    if action == "add":
        if not name:
            raise ValidationError("name required")
        wb.create_sheet(name[:31], index=index)
    elif action == "delete":
        if not name:
            raise ValidationError("name required")
        if name not in wb.sheetnames:
            raise NotFoundError(name)
        if len(wb.sheetnames) == 1:
            raise ValidationError("Cannot delete the last sheet")
        del wb[name]
    elif action == "rename":
        if not name or not new_name:
            raise ValidationError("name and new_name required")
        wb[name].title = new_name[:31]
    elif action == "activate":
        if not name:
            raise ValidationError("name required")
        wb.active = wb.sheetnames.index(name)
    else:
        raise ValidationError(f"Unknown sheet action: {action}")
    names = list(wb.sheetnames)
    wb.save(str(path))
    wb.close()
    return {"sheets": names, "action": action}


def format_range(path: Path, sheet: str | None, range_addr: str, spec: dict[str, Any]) -> dict[str, Any]:
    wb = load(path)
    ws = _sheet(wb, sheet)
    font_kw: dict[str, Any] = {}
    if spec.get("bold") is not None:
        font_kw["bold"] = bool(spec["bold"])
    if spec.get("italic") is not None:
        font_kw["italic"] = bool(spec["italic"])
    if spec.get("size"):
        font_kw["size"] = float(spec["size"])
    if spec.get("name"):
        font_kw["name"] = spec["name"]
    if spec.get("color"):
        font_kw["color"] = spec["color"].lstrip("#")
    fill = None
    if spec.get("fill"):
        fill = PatternFill("solid", fgColor=spec["fill"].lstrip("#"))
    number_format = spec.get("number_format")
    align = None
    if spec.get("align") or spec.get("valign") or spec.get("wrap"):
        align = Alignment(
            horizontal=spec.get("align"),
            vertical=spec.get("valign"),
            wrap_text=bool(spec.get("wrap")),
        )
    border = None
    if spec.get("border"):
        side = Side(style=spec.get("border_style") or "thin")
        border = Border(left=side, right=side, top=side, bottom=side)
    for row in ws[range_addr]:
        cells = row if isinstance(row, tuple) else (row,)
        for cell in cells:
            if font_kw:
                cell.font = Font(**font_kw)
            if fill is not None:
                cell.fill = fill
            if number_format:
                cell.number_format = number_format
            if align is not None:
                cell.alignment = align
            if border is not None:
                cell.border = border
    if spec.get("column_width"):
        min_col, _, max_col, _ = range_boundaries(range_addr)
        for c in range(min_col, max_col + 1):
            ws.column_dimensions[get_column_letter(c)].width = float(spec["column_width"])
    if spec.get("row_height"):
        _, min_row, _, max_row = range_boundaries(range_addr)
        for r in range(min_row, max_row + 1):
            ws.row_dimensions[r].height = float(spec["row_height"])
    if spec.get("color_scale"):
        ws.conditional_formatting.add(
            range_addr,
            ColorScaleRule(
                start_type="min",
                start_color="63BE7B",
                end_type="max",
                end_color="F8696B",
            ),
        )
    if spec.get("validation"):
        dv = DataValidation(type=spec["validation"].get("type", "list"), formula1=spec["validation"].get("formula1"))
        dv.add(range_addr)
        ws.add_data_validation(dv)
    wb.save(str(path))
    wb.close()
    return {"path": str(path), "range": range_addr}


def structure(path: Path, sheet: str | None, spec: dict[str, Any]) -> dict[str, Any]:
    wb = load(path)
    ws = _sheet(wb, sheet)
    info: dict[str, Any] = {}
    if spec.get("merge"):
        ws.merge_cells(spec["merge"])
        info["merged"] = spec["merge"]
    if spec.get("unmerge"):
        ws.unmerge_cells(spec["unmerge"])
        info["unmerged"] = spec["unmerge"]
    if spec.get("freeze"):
        ws.freeze_panes = spec["freeze"]
        info["freeze"] = spec["freeze"]
    if spec.get("auto_filter") is True:
        ws.auto_filter.ref = ws.dimensions
        info["auto_filter"] = ws.auto_filter.ref
    elif isinstance(spec.get("auto_filter"), str):
        ws.auto_filter.ref = spec["auto_filter"]
        info["auto_filter"] = spec["auto_filter"]
    if spec.get("table"):
        t = spec["table"]
        name = t.get("name") or "Table1"
        ref = t.get("ref") or ws.dimensions
        table = Table(displayName=name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name=t.get("style") or "TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)
        info["table"] = name
    if spec.get("defined_name"):
        dn = spec["defined_name"]
        defn = DefinedName(dn["name"], attr_text=dn["ref"])
        wb.defined_names.add(defn)
        info["defined_name"] = dn["name"]
    wb.save(str(path))
    wb.close()
    return {"path": str(path), **info}


def add_chart(path: Path, sheet: str | None, spec: dict[str, Any]) -> dict[str, Any]:
    kind = (spec.get("type") or "bar").lower()
    if kind not in CHARTS:
        raise ValidationError(f"Unsupported chart type {kind}", {"supported": list(CHARTS)})
    wb = load(path)
    ws = _sheet(wb, sheet)
    data_ref = spec.get("data")
    cats = spec.get("categories")
    if not data_ref:
        raise ValidationError("chart requires data range e.g. B1:B10")
    min_col, min_row, max_col, max_row = range_boundaries(data_ref)
    chart = CHARTS[kind]()
    chart.title = spec.get("title") or ""
    values = Reference(ws, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)
    if kind == "pie":
        chart.add_data(values, titles_from_data=bool(spec.get("titles_from_data", True)))
    else:
        chart.add_data(values, titles_from_data=bool(spec.get("titles_from_data", True)))
    if cats:
        cmin_col, cmin_row, cmax_col, cmax_row = range_boundaries(cats)
        chart.set_categories(Reference(ws, min_col=cmin_col, min_row=cmin_row, max_col=cmax_col, max_row=cmax_row))
    anchor = spec.get("anchor") or "E2"
    ws.add_chart(chart, anchor)
    wb.save(str(path))
    wb.close()
    return {"path": str(path), "type": kind, "anchor": anchor}


def analyze(path: Path, sheet: str | None = None, max_rows: int = 50_000) -> dict[str, Any]:
    try:
        import pandas as pd
    except ImportError as exc:  # pragma: no cover
        raise ValidationError("pandas is required for excel_analyze") from exc
    sheet_name = sheet or 0
    df = pd.read_excel(path, sheet_name=sheet_name, nrows=max_rows)
    if not isinstance(df, pd.DataFrame):
        # multiple sheets
        first = next(iter(df.values()))
        df = first
    profile = []
    for col in df.columns:
        s = df[col]
        non_null = s.dropna()
        sample = non_null.head(5).astype(str).tolist()
        profile.append(
            {
                "column": str(col),
                "dtype": str(s.dtype),
                "non_null": int(s.notna().sum()),
                "nulls": int(s.isna().sum()),
                "unique": int(s.nunique(dropna=True)),
                "sample": sample,
            }
        )
    numeric = df.select_dtypes(include="number")
    stats = {}
    if not numeric.empty:
        desc = numeric.describe().to_dict()
        stats = {str(k): {sk: _json_num(sv) for sk, sv in v.items()} for k, v in desc.items()}
    dup_rows = int(df.duplicated().sum())
    return {
        "sheet": sheet or (df.attrs.get("name") if hasattr(df, "attrs") else None),
        "rows": int(len(df)),
        "columns": [str(c) for c in df.columns],
        "profile": profile,
        "numeric_stats": stats,
        "duplicate_rows": dup_rows,
        "empty": bool(df.empty),
    }


def _json_num(v: Any) -> Any:
    try:
        if hasattr(v, "item"):
            v = v.item()
    except Exception:
        return str(v)
    if isinstance(v, float):
        return round(v, 6)
    return v


def to_csv(path: Path, dest: Path, sheet: str | None = None) -> None:
    import pandas as pd

    df = pd.read_excel(path, sheet_name=sheet or 0)
    if not isinstance(df, pd.DataFrame):
        df = next(iter(df.values()))
    df.to_csv(dest, index=False)


def from_csv(csv_path: Path, dest: Path, sheet: str = "Sheet1") -> None:
    import pandas as pd

    df = pd.read_csv(csv_path)
    df.to_excel(dest, sheet_name=sheet[:31], index=False)


def get_props(path: Path) -> dict[str, Any]:
    wb = load(path, read_only=True)
    try:
        p = wb.properties
        return {
            "title": p.title,
            "creator": p.creator,
            "subject": p.subject,
            "keywords": p.keywords,
            "created": str(p.created) if p.created else None,
            "modified": str(p.modified) if p.modified else None,
        }
    finally:
        wb.close()


def set_props(path: Path, props: dict[str, Any]) -> dict[str, Any]:
    wb = load(path)
    p = wb.properties
    for k in ("title", "creator", "subject", "keywords", "description"):
        if k in props and props[k] is not None:
            setattr(p, k, str(props[k]))
    wb.save(str(path))
    wb.close()
    return get_props(path)
