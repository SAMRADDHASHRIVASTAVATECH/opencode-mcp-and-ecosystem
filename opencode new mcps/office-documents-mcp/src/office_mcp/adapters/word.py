"""Word (DOCX) adapter over python-docx."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from office_mcp.errors import UnsupportedError, ValidationError

ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def create_document(path: Path, title: str | None = None) -> None:
    doc = Document()
    if title:
        doc.core_properties.title = title
        doc.add_heading(title, level=0)
    doc.save(str(path))


def open_doc(path: Path) -> Document:
    return Document(str(path))


def extract(path: Path, include_tables: bool = True, include_headers: bool = True) -> dict[str, Any]:
    doc = open_doc(path)
    paragraphs = []
    for p in doc.paragraphs:
        paragraphs.append(
            {
                "text": p.text,
                "style": p.style.name if p.style is not None else None,
                "align": str(p.alignment) if p.alignment is not None else None,
            }
        )
    tables = []
    if include_tables:
        for table in doc.tables:
            rows = []
            for row in table.rows:
                rows.append([cell.text for cell in row.cells])
            tables.append({"rows": rows, "n_rows": len(rows), "n_cols": len(rows[0]) if rows else 0})
    headers, footers = [], []
    if include_headers:
        for i, section in enumerate(doc.sections):
            headers.append({"section": i, "text": section.header.paragraphs[0].text if section.header.paragraphs else ""})
            footers.append({"section": i, "text": section.footer.paragraphs[0].text if section.footer.paragraphs else ""})
    comments = _comments(doc)
    images = []
    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            images.append({"content_type": rel.target_ref, "reltype": rel.reltype})
    styles = [s.name for s in doc.styles if s.name]
    plain = "\n".join(p["text"] for p in paragraphs if p["text"])
    return {
        "paragraphs": paragraphs,
        "tables": tables,
        "headers": headers,
        "footers": footers,
        "comments": comments,
        "images": images,
        "styles": styles[:80],
        "sections": len(doc.sections),
        "plain_text": plain,
        "paragraph_count": len(paragraphs),
        "table_count": len(tables),
    }


def _comments(doc: Document) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    comments_part = getattr(doc.part, "comments_part", None)
    if comments_part is None:
        return out
    comments = getattr(comments_part, "comments", None)
    if comments is None:
        return out
    try:
        for c in comments:
            text = getattr(c, "text", None) or ""
            author = getattr(c, "author", None)
            out.append({"author": author, "text": text})
    except Exception:
        return out
    return out


def apply_ops(path: Path, ops: list[dict[str, Any]]) -> dict[str, Any]:
    doc = open_doc(path)
    applied = 0
    for op in ops:
        if not isinstance(op, dict) or "op" not in op:
            raise ValidationError("Each edit op must be an object with 'op'")
        kind = op["op"]
        if kind == "add_heading":
            doc.add_heading(op.get("text") or "", level=int(op.get("level", 1)))
        elif kind == "add_paragraph":
            p = doc.add_paragraph()
            _add_runs(p, op)
            if op.get("style"):
                p.style = op["style"]
            if op.get("align") in ALIGN:
                p.alignment = ALIGN[op["align"]]
        elif kind == "add_list":
            style = "List Number" if op.get("ordered") else "List Bullet"
            for item in op.get("items") or []:
                doc.add_paragraph(str(item), style=style)
        elif kind == "add_table":
            rows = op.get("rows") or []
            if not rows:
                raise ValidationError("add_table requires rows")
            ncols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=ncols)
            table.style = op.get("style") or "Table Grid"
            for i, row in enumerate(rows):
                for j, val in enumerate(row):
                    table.rows[i].cells[j].text = "" if val is None else str(val)
        elif kind == "add_page_break":
            doc.add_page_break()
        elif kind == "add_picture":
            pic = op.get("path")
            if not pic:
                raise ValidationError("add_picture requires path")
            width = Inches(float(op["width_in"])) if op.get("width_in") else None
            doc.add_picture(pic, width=width)
        elif kind == "add_comment":
            _add_comment(doc, op)
        else:
            raise UnsupportedError(f"Unknown word op: {kind}")
        applied += 1
    doc.save(str(path))
    return {"applied": applied, "path": str(path)}


def _add_runs(paragraph: Any, op: dict[str, Any]) -> None:
    runs = op.get("runs")
    if runs:
        for r in runs:
            run = paragraph.add_run(r.get("text") or "")
            run.bold = bool(r.get("bold"))
            run.italic = bool(r.get("italic"))
            run.underline = bool(r.get("underline"))
            if r.get("size_pt"):
                run.font.size = Pt(float(r["size_pt"]))
            if r.get("font"):
                run.font.name = r["font"]
            if r.get("color"):
                run.font.color.rgb = RGBColor.from_string(r["color"].lstrip("#"))
        return
    run = paragraph.add_run(op.get("text") or "")
    run.bold = bool(op.get("bold"))
    run.italic = bool(op.get("italic"))
    if op.get("size_pt"):
        run.font.size = Pt(float(op["size_pt"]))
    if op.get("font"):
        run.font.name = op["font"]


def _add_comment(doc: Document, op: dict[str, Any]) -> None:
    text = op.get("on_text") or ""
    comment = op.get("comment") or ""
    author = op.get("author") or "office-mcp"
    target_run = None
    for p in doc.paragraphs:
        for run in p.runs:
            if text and text in (run.text or ""):
                target_run = run
                break
        if target_run:
            break
    add = getattr(doc, "add_comment", None)
    if add is None:
        raise UnsupportedError("This python-docx build cannot add comments; upgrade to >= 1.2")
    if target_run is None:
        p = doc.add_paragraph(text or comment)
        target_run = p.runs[0] if p.runs else p.add_run(text or ".")
    add(target_run, comment, author=author)


def find_replace(path: Path, find: str, replace: str, regex: bool = False, count: int = 0) -> dict[str, Any]:
    doc = open_doc(path)
    n = 0
    limit = count if count and count > 0 else 10_000

    def sub(text: str) -> str:
        nonlocal n
        if n >= limit:
            return text
        if regex:
            def repl(m: re.Match[str]) -> str:
                nonlocal n
                if n >= limit:
                    return m.group(0)
                n += 1
                return replace
            return re.sub(find, repl, text)
        if find not in text:
            return text
        pieces = text.split(find)
        out = pieces[0]
        for piece in pieces[1:]:
            if n < limit:
                out += replace
                n += 1
            else:
                out += find
            out += piece
        return out

    for p in doc.paragraphs:
        if n >= limit:
            break
        full = p.text
        new = sub(full)
        if new != full:
            if p.runs:
                p.runs[0].text = new
                for r in p.runs[1:]:
                    r.text = ""
            else:
                p.add_run(new)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    full = p.text
                    new = sub(full)
                    if new != full:
                        if p.runs:
                            p.runs[0].text = new
                            for r in p.runs[1:]:
                                r.text = ""
                        else:
                            p.add_run(new)
    doc.save(str(path))
    return {"replacements": n, "path": str(path)}


def layout(path: Path, spec: dict[str, Any]) -> dict[str, Any]:
    doc = open_doc(path)
    section = doc.sections[int(spec.get("section", 0))]
    if spec.get("header") is not None:
        hp = section.header.paragraphs[0] if section.header.paragraphs else section.header.add_paragraph()
        hp.text = str(spec["header"])
    if spec.get("footer") is not None:
        fp = section.footer.paragraphs[0] if section.footer.paragraphs else section.footer.add_paragraph()
        fp.text = str(spec["footer"])
    if spec.get("page_width_in"):
        section.page_width = Inches(float(spec["page_width_in"]))
    if spec.get("page_height_in"):
        section.page_height = Inches(float(spec["page_height_in"]))
    if spec.get("orientation") == "landscape":
        section.page_width, section.page_height = section.page_height, section.page_width
    if spec.get("margins_in"):
        m = spec["margins_in"]
        if "top" in m:
            section.top_margin = Inches(float(m["top"]))
        if "bottom" in m:
            section.bottom_margin = Inches(float(m["bottom"]))
        if "left" in m:
            section.left_margin = Inches(float(m["left"]))
        if "right" in m:
            section.right_margin = Inches(float(m["right"]))
    doc.save(str(path))
    return {"path": str(path), "sections": len(doc.sections)}


def markdown_to_docx(md: str, path: Path, title: str | None = None) -> None:
    doc = Document()
    if title:
        doc.core_properties.title = title
        doc.add_heading(title, level=0)
    lines = md.replace("\r\n", "\n").split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*[-: ]+\|", lines[i + 1]):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                row = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not re.match(r"^[-: ]+$", "".join(row).replace("|", "")):
                    if not all(re.match(r"^:?-+:?$", c) for c in row):
                        rows.append(row)
                i += 1
            if rows:
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = "Table Grid"
                for r_i, row in enumerate(rows):
                    for c_i, val in enumerate(row):
                        if c_i < len(table.rows[r_i].cells):
                            table.rows[r_i].cells[c_i].text = val
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            doc.add_heading(heading.group(2), level=len(heading.group(1)))
            i += 1
            continue
        if re.match(r"^[-*]\s+", line):
            while i < len(lines) and re.match(r"^[-*]\s+", lines[i]):
                doc.add_paragraph(re.sub(r"^[-*]\s+", "", lines[i]), style="List Bullet")
                i += 1
            continue
        if re.match(r"^\d+\.\s+", line):
            while i < len(lines) and re.match(r"^\d+\.\s+", lines[i]):
                doc.add_paragraph(re.sub(r"^\d+\.\s+", "", lines[i]), style="List Number")
                i += 1
            continue
        p = doc.add_paragraph()
        _inline_md(p, line)
        i += 1
    doc.save(str(path))


def _inline_md(paragraph: Any, text: str) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Courier New"
        else:
            paragraph.add_run(part)


def set_core_props(path: Path, props: dict[str, Any]) -> dict[str, str]:
    doc = open_doc(path)
    cp = doc.core_properties
    mapping = {
        "title": "title",
        "subject": "subject",
        "author": "author",
        "keywords": "keywords",
        "category": "category",
        "comments": "comments",
    }
    for k, attr in mapping.items():
        if k in props and props[k] is not None:
            setattr(cp, attr, str(props[k]))
    doc.save(str(path))
    return get_core_props(path)


def get_core_props(path: Path) -> dict[str, Any]:
    doc = open_doc(path)
    cp = doc.core_properties
    return {
        "title": cp.title,
        "subject": cp.subject,
        "author": cp.author,
        "keywords": cp.keywords,
        "category": cp.category,
        "comments": cp.comments,
        "created": str(cp.created) if cp.created else None,
        "modified": str(cp.modified) if cp.modified else None,
        "last_modified_by": cp.last_modified_by,
    }


# silence unused qn import if styles need it later
_ = qn
