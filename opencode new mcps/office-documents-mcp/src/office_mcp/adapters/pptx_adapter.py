"""PowerPoint (PPTX) adapter over python-pptx."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

from office_mcp.errors import NotFoundError, ValidationError

ALIGN = {
    "left": PP_ALIGN.LEFT,
    "center": PP_ALIGN.CENTER,
    "right": PP_ALIGN.RIGHT,
}


def create(path: Path, title: str | None = None, blank: bool = False) -> None:
    prs = Presentation()
    if not blank:
        layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(layout)
        if slide.shapes.title is not None:
            slide.shapes.title.text = title or "Presentation"
        for shape in slide.placeholders:
            if shape.has_text_frame and shape != slide.shapes.title:
                shape.text_frame.text = ""
    elif title:
        prs.core_properties.title = title
    if title:
        prs.core_properties.title = title
    prs.save(str(path))


def open_prs(path: Path) -> Presentation:
    return Presentation(str(path))


def extract(path: Path) -> dict[str, Any]:
    prs = open_prs(path)
    slides = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        notes = ""
        shapes = []
        for shape in slide.shapes:
            info: dict[str, Any] = {
                "name": shape.name,
                "shape_type": str(getattr(shape, "shape_type", None)),
                "has_text": bool(getattr(shape, "has_text_frame", False)),
            }
            if getattr(shape, "has_text_frame", False):
                t = shape.text_frame.text
                info["text"] = t
                texts.append(t)
            if getattr(shape, "has_table", False):
                table = shape.table
                rows = [[cell.text for cell in row.cells] for row in table.rows]
                info["table"] = rows
            shapes.append(info)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text
        layout_name = None
        try:
            layout_name = slide.slide_layout.name
        except Exception:
            pass
        slides.append(
            {
                "index": i,
                "layout": layout_name,
                "texts": texts,
                "notes": notes,
                "shapes": shapes,
            }
        )
    return {
        "slide_count": len(slides),
        "slides": slides,
        "slide_width": int(prs.slide_width),
        "slide_height": int(prs.slide_height),
        "plain_text": "\n\n".join(
            f"Slide {s['index']}\n" + "\n".join(x for x in s["texts"] if x) for s in slides
        ),
        "core": {
            "title": prs.core_properties.title,
            "author": prs.core_properties.author,
            "subject": prs.core_properties.subject,
        },
    }


def apply_ops(path: Path, ops: list[dict[str, Any]]) -> dict[str, Any]:
    prs = open_prs(path)
    applied = 0
    for op in ops:
        kind = op.get("op")
        if kind == "add_slide":
            layout_idx = int(op.get("layout", 1))
            if layout_idx < 0 or layout_idx >= len(prs.slide_layouts):
                layout_idx = 1 if len(prs.slide_layouts) > 1 else 0
            slide = prs.slides.add_slide(prs.slide_layouts[layout_idx])
            if op.get("title") and slide.shapes.title is not None:
                slide.shapes.title.text = op["title"]
            body = op.get("body")
            if body:
                _set_body(slide, body)
        elif kind == "set_title":
            slide = _slide(prs, op.get("slide", 1))
            if slide.shapes.title is None:
                raise ValidationError("Slide has no title placeholder")
            slide.shapes.title.text = op.get("text") or ""
        elif kind == "set_body":
            slide = _slide(prs, op.get("slide", 1))
            _set_body(slide, op.get("text") or op.get("body") or "")
        elif kind == "add_textbox":
            slide = _slide(prs, op.get("slide", 1))
            left = Inches(float(op.get("left_in", 0.5)))
            top = Inches(float(op.get("top_in", 2)))
            width = Inches(float(op.get("width_in", 9)))
            height = Inches(float(op.get("height_in", 1)))
            box = slide.shapes.add_textbox(left, top, width, height)
            tf = box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = op.get("text") or ""
            if op.get("align") in ALIGN:
                p.alignment = ALIGN[op["align"]]
            if op.get("size_pt"):
                p.font.size = Pt(float(op["size_pt"]))
            if op.get("bold"):
                p.font.bold = True
            if op.get("color"):
                p.font.color.rgb = RGBColor.from_string(op["color"].lstrip("#"))
        elif kind == "add_picture":
            slide = _slide(prs, op.get("slide", 1))
            slide.shapes.add_picture(
                op["path"],
                Inches(float(op.get("left_in", 1))),
                Inches(float(op.get("top_in", 1))),
                width=Inches(float(op["width_in"])) if op.get("width_in") else None,
            )
        elif kind == "add_table":
            slide = _slide(prs, op.get("slide", 1))
            rows = op.get("rows") or []
            if not rows:
                raise ValidationError("add_table requires rows")
            n_rows, n_cols = len(rows), max(len(r) for r in rows)
            table_shape = slide.shapes.add_table(
                n_rows,
                n_cols,
                Inches(float(op.get("left_in", 0.5))),
                Inches(float(op.get("top_in", 2))),
                Inches(float(op.get("width_in", 9))),
                Inches(float(op.get("height_in", 0.4 * n_rows))),
            )
            tbl = table_shape.table
            for r_i, row in enumerate(rows):
                for c_i, val in enumerate(row):
                    tbl.cell(r_i, c_i).text = "" if val is None else str(val)
        elif kind == "set_notes":
            slide = _slide(prs, op.get("slide", 1))
            slide.notes_slide.notes_text_frame.text = op.get("text") or ""
        elif kind == "delete_slide":
            idx = int(op.get("slide", 1)) - 1
            _delete_slide(prs, idx)
        else:
            raise ValidationError(f"Unknown pptx op: {kind}")
        applied += 1
    prs.save(str(path))
    return {"applied": applied, "slide_count": len(prs.slides), "path": str(path)}


def _slide(prs: Presentation, n: int):
    idx = int(n) - 1
    if idx < 0 or idx >= len(prs.slides):
        raise NotFoundError(f"Slide {n} not found", {"slide_count": len(prs.slides)})
    return prs.slides[idx]


def _set_body(slide: Any, text: str) -> None:
    lines = text.split("\n") if isinstance(text, str) else list(text)
    body_shape = None
    for shape in slide.placeholders:
        if shape.has_text_frame and shape != slide.shapes.title:
            body_shape = shape
            break
    if body_shape is None:
        for shape in slide.shapes:
            if shape.has_text_frame and shape != slide.shapes.title:
                body_shape = shape
                break
    if body_shape is None:
        raise ValidationError("No body placeholder on this slide; use add_textbox")
    tf = body_shape.text_frame
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0


def _delete_slide(prs: Presentation, index: int) -> None:
    if index < 0 or index >= len(prs.slides):
        raise NotFoundError(f"Slide {index + 1} not found")
    rId = prs.slides._sldIdLst[index].rId  # noqa: SLF001
    prs.part.drop_rel(rId)
    sldId = prs.slides._sldIdLst[index]  # noqa: SLF001
    prs.slides._sldIdLst.remove(sldId)  # noqa: SLF001


def get_props(path: Path) -> dict[str, Any]:
    prs = open_prs(path)
    cp = prs.core_properties
    return {
        "title": cp.title,
        "author": cp.author,
        "subject": cp.subject,
        "keywords": cp.keywords,
        "created": str(cp.created) if cp.created else None,
        "modified": str(cp.modified) if cp.modified else None,
    }


def set_props(path: Path, props: dict[str, Any]) -> dict[str, Any]:
    prs = open_prs(path)
    cp = prs.core_properties
    for k in ("title", "author", "subject", "keywords", "comments"):
        if k in props and props[k] is not None:
            setattr(cp, k, str(props[k]))
    prs.save(str(path))
    return get_props(path)


_ = (Emu,)
